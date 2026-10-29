"""Produce the Italian release dossier from versioned sources and recorded evidence."""
from pathlib import Path
import hashlib
import json
import re
import shutil
import sys
import textwrap
import zipfile
import xml.etree.ElementTree as ET
from datetime import datetime

ROOT = Path(__file__).resolve().parents[1]
# Optional workspace-local documentation dependencies; never part of application runtime.
LOCAL = ROOT / '.test-results/python'
if LOCAL.exists():
    sys.path.insert(0, str(LOCAL))
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw, ImageFont

DOCS = ROOT / 'docs'
EVIDENCE = DOCS / 'evidenze'
SOURCES = DOCS / 'sorgenti'
VERSION = json.loads((ROOT / 'package.json').read_text(encoding='utf-8'))['version']
RED = 'AA1616'


def capture():
    source = ROOT / '.test-results/release/results.json'
    data = json.loads(source.read_text(encoding='utf-8'))
    if len(data['results']) != 12 or any(row['result'] != 'PASS' for row in data['results']):
        raise RuntimeError('La verifica completa deve terminare prima della cattura delle evidenze.')
    EVIDENCE.mkdir(exist_ok=True)
    for row in data['results']:
        src = ROOT / row['log']
        content = src.read_bytes()
        if hashlib.sha256(content).hexdigest() != row['sha256']:
            raise RuntimeError('Evidenza modificata: ' + row['name'])
        dest = EVIDENCE / (row['name'] + '.log')
        shutil.copyfile(src, dest)
        row['log'] = 'evidenze/' + dest.name
    browser = ROOT / '.test-results/browser/report.json'
    data['browser'] = json.loads(browser.read_text(encoding='utf-8'))
    for src in (ROOT / '.test-results/browser').glob('*.png'):
        shutil.copyfile(src, EVIDENCE / src.name)
    data['version'] = VERSION
    data['scope'] = 'Windows, PostgreSQL 18.6 locale isolato, provider Stripe/Resend controllati, browser Chromium e mobile emulato. Nessuna vendita remota.'
    (EVIDENCE / 'esecuzione.json').write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')


def draw_flow(name, title, labels, note=''):
    width, box_height, gap = 1500, 95, 45
    height = 150 + len(labels) * (box_height + gap) + 65
    img = Image.new('RGB', (width, height), '#ffffff')
    canvas = ImageDraw.Draw(img)
    font_path = Path('C:/Windows/Fonts/arial.ttf')
    font = ImageFont.truetype(str(font_path), 27) if font_path.exists() else ImageFont.load_default(size=27)
    title_font = ImageFont.truetype(str(font_path), 38) if font_path.exists() else ImageFont.load_default(size=38)
    canvas.text((70, 35), title, fill='#aa1616', font=title_font)
    for index, label in enumerate(labels):
        top = 120 + index * (box_height + gap)
        canvas.rounded_rectangle((80, top, width - 80, top + box_height), radius=12, fill='#f5f5f5', outline='#aa1616', width=3)
        canvas.multiline_text((105, top + 16), '\n'.join(textwrap.wrap(label, 90)), font=font, fill='#161616', spacing=5)
        if index < len(labels) - 1:
            x, y = width // 2, top + box_height
            canvas.line((x, y, x, y + gap - 8), fill='#aa1616', width=4)
            canvas.polygon([(x - 9, y + gap - 18), (x + 9, y + gap - 18), (x, y + gap - 4)], fill='#aa1616')
    if note:
        canvas.text((80, height - 55), note, font=font, fill='#555555')
    img.save(EVIDENCE / (name + '.png'))


def relation_diagram(name, title, nodes, edges, note):
    img = Image.new('RGB', (1600, 1100), '#ffffff')
    canvas = ImageDraw.Draw(img)
    font_path = Path('C:/Windows/Fonts/arial.ttf')
    font = ImageFont.truetype(str(font_path), 23) if font_path.exists() else ImageFont.load_default(size=23)
    heading = ImageFont.truetype(str(font_path), 35) if font_path.exists() else ImageFont.load_default(size=35)
    canvas.text((55, 35), title, font=heading, fill='#aa1616')
    for start, end, label, logical in edges:
        x1, y1, _, _ = nodes[start]
        x2, y2, _, _ = nodes[end]
        a, b = (x1 + 210, y1 + 75), (x2 + 210, y2 + 75)
        color = '#777777' if logical else '#aa1616'
        canvas.line((a, b), fill=color, width=4)
        center = ((a[0] + b[0]) / 2, (a[1] + b[1]) / 2)
        bounds = canvas.textbbox(center, label, font=font, anchor='mm')
        canvas.rectangle((bounds[0]-8, bounds[1]-5, bounds[2]+8, bounds[3]+5), fill='white')
        canvas.text(center, label, font=font, fill=color, anchor='mm')
    for x, y, heading_text, body in nodes.values():
        canvas.rounded_rectangle((x, y, x + 420, y + 150), radius=12, fill='#f5f5f5', outline='#aa1616', width=3)
        canvas.text((x + 18, y + 17), heading_text, font=font, fill='#aa1616')
        canvas.multiline_text((x + 18, y + 57), body, font=font, fill='#161616', spacing=7)
    canvas.multiline_text((55, 995), '\n'.join(textwrap.wrap(note, 110)), font=font, fill='#555555', spacing=8)
    img.save(EVIDENCE / (name + '.png'))


def diagrams():
    draw_flow('flusso-galleria', 'Galleria prodotto — FUD-001', [
        'Fotografie: fronte, retro, dettaglio. Taglia e quantità restano in Storefront.',
        'Selezione 3D → import dinamico del viewer → download GLB con texture incorporate.',
        'Caricamento riuscito → controlli attivi. Errore o timeout → fallback fotografico.',
        'Mouse / touch / zoom controllano il modello senza cambiare slide.',
        'Selettore foto → smontaggio viewer → fotografie, senza ricaricare la pagina.'
    ])
    draw_flow('flusso-ordine', 'Ordine e conferma — FUD-002', [
        'Cliente → /api/status → taglia, quantità, dati, condizioni e verifica anti-bot.',
        'POST /api/orders → Origin, rate limit, Zod e Turnstile.',
        'Transazione PostgreSQL → lock Settings → controlli → Order con snapshot.',
        'Stripe Checkout → pagamento → webhook firmato /api/stripe/webhook.',
        'Verifica modalità/importi/identificativi → transazione idempotente → PAID.',
        'Outbox conferma + controllo soglia → /conferma e contatore aggiornato.'
    ])
    draw_flow('flusso-admin', 'Identità amministrativa — FUD-005', [
        'Gestore → migrazione additiva → create-admin → AdminUser ADMIN (solo hash).',
        'Login → Origin + budget DB + Turnstile → account attivo + verifica bcrypt.',
        'JWT firmato + digest jti e impronta credenziale nel DB → cookie protetto.',
        'Pagine/API → firma + scadenza + revoca + account attivo + impronta corrente.',
        'Logout → revoca DB → cookie cancellato. Errore DB → accesso negato.'
    ])
    relation_diagram('architettura', 'RUSH Presale — connessioni fra componenti', {
        'browser': (55, 130, 'Browser cliente / admin', 'React, foto, model-viewer\nHTTP e asset locali'),
        'next': (590, 430, 'Next.js su Vercel', 'Proxy, API, servizi\nPrisma e outbox'),
        'db': (1125, 130, 'PostgreSQL', 'Ordini, impostazioni\nAccount e sessioni'),
        'stripe': (55, 730, 'Stripe', 'Checkout e rimborsi\nWebhook firmati'),
        'providers': (1125, 730, 'Resend / Turnstile / Redis', 'Email, anti-bot\nLimite distribuito'),
        'cron': (590, 130, 'Vercel Cron', 'daily / cleanup\nBearer CRON_SECRET'),
    }, [('browser', 'next', 'HTTP', False), ('next', 'db', 'SQL / TLS', False),
        ('stripe', 'next', 'API / webhook', False), ('next', 'providers', 'HTTPS', False),
        ('cron', 'next', 'HTTP autenticato', False)],
        'Ogni linea indica una comunicazione diretta. Gli asset GLB e fotografici sono serviti dalla stessa applicazione.')
    relation_diagram('modello-dati', 'Modello entità-relazioni — nucleo operativo', {
        'terms': (590, 130, 'TermsVersion', 'PK id\nversion, text'),
        'settings': (55, 130, 'Settings', 'PK id = global\nFK currentTermsId'),
        'order': (590, 430, 'Order', 'PK id\nFK termsVersionId'),
        'refund': (1125, 430, 'RefundOperation', 'PK id\nFK orderId UNIQUE'),
        'admin': (55, 730, 'AdminUser', 'PK id, email UNIQUE\nrole, active, passwordHash'),
        'session': (1125, 730, 'AdminSession', 'PK id (digest jti)\nactor, credentialHash'),
    }, [('settings', 'terms', 'N : 0..1', False), ('order', 'terms', 'N : 0..1', False),
        ('order', 'refund', '1 : 0..1', False), ('admin', 'session', '1 : N, logico', True)],
        'Rosso: foreign key; grigio: riferimento applicativo senza foreign key. Outbox, ledger webhook, audit, rate limit ed errori sono descritti nel dizionario dati; non si inventano relazioni SQL.')


def clean(text):
    return re.sub(r'[*`]', '', re.sub(r'\[([^]]+)\]\(([^)]+)\)', r'\1 (\2)', text))


def field(paragraph, instruction):
    field_element = OxmlElement('w:fldSimple')
    field_element.set(qn('w:instr'), instruction)
    paragraph._p.append(field_element)


def make_docx(source, extra=''):
    content = source.read_text(encoding='utf-8') + extra
    lines = content.splitlines()
    title = lines[0].removeprefix('# ')
    doc = Document()
    section = doc.sections[0]
    section.page_height, section.page_width = Cm(29.7), Cm(21)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(2)
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = 'Calibri', Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    for level in range(1, 4):
        doc.styles[f'Heading {level}'].font.color.rgb = RGBColor.from_string(RED)
    section.header.paragraphs[0].text = 'RUSH  /  PRE-SALE                                         DOSSIER TECNICO  ' + VERSION
    section.header.paragraphs[0].style = doc.styles['Caption']
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('RUSH Presale · ' + DATE + '  |  Pagina ')
    field(footer, 'PAGE')
    footer.add_run(' di ')
    field(footer, 'NUMPAGES')
    doc.add_paragraph('RUSH', 'Title')
    doc.add_paragraph(title, 'Subtitle')
    doc.add_paragraph('Versione documento 1.0 · Candidato applicativo ' + VERSION)
    doc.add_paragraph('Data: ' + DATE + ' · Lingua: italiano')
    doc.add_paragraph('Stato: consegna tecnica del branch feature/rush-final-release.\nAccettazione commerciale subordinata alle verifiche remote indicate nel rapporto.')
    doc.add_page_break()
    doc.add_heading('Indice', level=1)
    for line in lines:
        if line.startswith('## '):
            doc.add_paragraph(clean(line[3:]))
    doc.add_page_break()
    index = 1
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith('|'):
            rows = []
            while index < len(lines) and lines[index].strip().startswith('|'):
                values = [clean(value.strip()) for value in lines[index].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'[:\-\s]+', value or '-') for value in values):
                    rows.append(values)
                index += 1
            table = doc.add_table(rows=0, cols=max(map(len, rows)))
            table.style = 'Light Shading Accent 1'
            for row_index, values in enumerate(rows):
                cells = table.add_row().cells
                for i, value in enumerate(values):
                    cells[i].text = value
                    for paragraph in cells[i].paragraphs:
                        for run in paragraph.runs:
                            run.font.size = Pt(8)
                            run.bold = row_index == 0
                if row_index == 0:
                    repeat = OxmlElement('w:tblHeader')
                    table.rows[0]._tr.get_or_add_trPr().append(repeat)
                cant_split = OxmlElement('w:cantSplit')
                table.rows[-1]._tr.get_or_add_trPr().append(cant_split)
            continue
        match = re.match(r'!\[([^]]+)\]\(([^)]+)\)', line)
        if match:
            path = (source.parent / match[2]).resolve()
            if not path.is_file():
                raise RuntimeError('Immagine mancante: ' + str(path))
            with Image.open(path) as img:
                w, h = img.size
            max_width, max_height = 16.5, 19
            width = min(max_width, max_height * w / h)
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run().add_picture(str(path), width=Cm(width))
            doc.add_paragraph(match[1], 'Caption')
        elif line.startswith('### '):
            doc.add_heading(clean(line[4:]), level=2)
        elif line.startswith('## '):
            doc.add_heading(clean(line[3:]), level=1)
        elif line.startswith('```'):
            index += 1
            while index < len(lines) and not lines[index].startswith('```'):
                p = doc.add_paragraph(lines[index])
                for run in p.runs:
                    run.font.name, run.font.size = 'Consolas', Pt(8)
                index += 1
        else:
            doc.add_paragraph(clean(line))
        index += 1
    doc.core_properties.title = title
    doc.core_properties.subject = 'RUSH Presale — consegna ' + VERSION
    doc.core_properties.author = 'RUSH Presale'
    path = DOCS / (source.stem + '.docx')
    doc.save(path)
    return path


def technical_appendix():
    lock = json.loads((ROOT / 'package-lock.json').read_text(encoding='utf-8'))
    package = json.loads((ROOT / 'package.json').read_text(encoding='utf-8'))
    text = '\n\n## 17. Versioni risolte delle dipendenze dirette\n\n| Pacchetto | Versione lockfile | Ambito |\n|---|---|---|\n'
    for group in ['dependencies', 'devDependencies']:
        for name in package[group]:
            version = lock['packages']['node_modules/' + name]['version']
            text += f'| {name} | {version} | {"runtime" if group == "dependencies" else "sviluppo"} |\n'
    text += '\n## 18. Dizionario del modello dati\n'
    schema = (ROOT / 'prisma/schema.prisma').read_text(encoding='utf-8')
    for name, body in re.findall(r'model (\w+) \{(.*?)\n\}', schema, re.S):
        text += '\n### ' + name + (' — storico ignorato' if '@@ignore' in body else '') + '\n\n| Campo | Tipo | Vincoli / attributi Prisma |\n|---|---|---|\n'
        for line in body.splitlines():
            parts = line.strip().split(maxsplit=2)
            if len(parts) >= 2 and not parts[0].startswith(('@@', '//')):
                text += '| ' + ' | '.join([parts[0], parts[1], parts[2] if len(parts) > 2 else '—']) + ' |\n'
    api = (DOCS / '04_DATI_API.md').read_text(encoding='utf-8')
    api = api[api.index('## API pubbliche'):]
    api = re.sub(r'^## (.+)$', r'### \1', api, flags=re.M)
    text += '\n## 19. Contratti HTTP\n' + api
    ops = (DOCS / '09_GESTIONE_TECNICA.md').read_text(encoding='utf-8')
    variables = ops[ops.index('| Gruppo |'):ops.index('Usare il generatore interattivo')]
    text += '\n## 20. Catalogo variabili\n' + variables
    return text


def requirement(label):
    low = label.lower()
    if any(x in low for x in ['home logo', 'home branding']):
        return 'BUR-001 / FUD-001'
    if any(x in low for x in ['backup', 'restore', 'migration']):
        return 'BUR-023, BUR-026 / FUD-010'
    if any(x in low for x in ['settings', 'configuration revision', 'impostazioni']):
        return 'BUR-021 / FUD-007'
    if any(x in low for x in ['3d', 'glb', 'webgl', 'texture', 'rotazione', 'carosello', 'foto']):
        return 'BUR-003, BUR-004, BUR-005 / FUD-001'
    if any(x in low for x in ['admin', 'session', 'login', 'logout', 'password', 'bcrypt', 'bootstrap', 'credential']):
        return 'BUR-015, BUR-016 / FUD-005'
    if any(x in low for x in ['refund', 'expiry', 'scaden', 'calendar']):
        return 'BUR-012, BUR-013 / FUD-004'
    if any(x in low for x in ['threshold', 'production', 'presale', 'soglia']):
        return 'BUR-006, BUR-007, BUR-008, BUR-014 / FUD-003'
    if any(x in low for x in ['csv', 'xlsx', 'export', 'dashboard', 'delivery', 'filter', 'pagination']):
        return 'BUR-017, BUR-018, BUR-020 / FUD-006'
    if any(x in low for x in ['email', 'mail', 'outbox']):
        return 'BUR-010, BUR-019 / FUD-008'
    if any(x in low for x in ['checkout', 'paid', 'payment', 'order', 'webhook', 'terms', 'quantity']):
        return 'BUR-002, BUR-009, BUR-010, BUR-011 / FUD-002'
    return 'BUR-022, BUR-026 / FUD-009, FUD-010'


def cases(data):
    rows = []
    tool_results = {
        'install': ('Installazione esatta dal lockfile e generazione del client Prisma.', 'Dipendenze installate; postinstall Prisma completato.'),
        'audit': ('Nessuna vulnerabilità npm di livello high o critical.', 'npm audit riporta zero vulnerabilità.'),
        'lint': ('Nessun errore o warning ESLint.', 'ESLint completato con --max-warnings=0.'),
        'typecheck': ('Nessun errore di tipo TypeScript.', 'tsc --noEmit completato senza diagnostiche.'),
        'db-validate': ('Schema Prisma sintatticamente e strutturalmente valido.', 'Prisma validate conferma la validità dello schema.'),
        'build': ('Client Prisma e applicazione Next compilabili per produzione.', 'Build completata; pagine, API e proxy elencati nel log.'),
        'test-migration': ('Migrazioni additive applicabili senza perdere ordini, importi, date o vincoli.', 'Migrazioni applicate e asserzioni di conservazione dati soddisfatte.'),
        'test-backup': ('Dump ripristinabile su DB separato mantenendo dati e checksum; sorgente intatta.', 'pg_dump/pg_restore riusciti; dati, relazioni e checksum verificati; sourceUnaffected=true.'),
    }
    for result in data['results']:
        log = (DOCS / result['log']).read_text(encoding='utf-8')
        log = re.sub(r'\x1b\[[0-9;]*m', '', log)
        labels = []
        if result['name'] in ['test', 'test-integration']:
            labels = re.findall(r'^\s*[✔✓] (.+?) \([\d.]+ms\)\s*$', log, re.M)
            labels = [label for label in labels if not label.startswith('PostgreSQL integration (')]
        elif result['name'] in ['test-http', 'test-browser']:
            labels = re.findall(r'^PASS (.+)$', log, re.M)
        if not labels:
            labels = [result['command']]
        for label in labels:
            expected, observed = tool_results.get(result['name'], (
                'Comportamento previsto dal caso: ' + label + '.',
                'PASS osservato nel log per il caso «' + label + '».'
            ))
            rows.append([f'TC-{len(rows)+1:03}', requirement(label), result['name'], label,
                         'Dipendenze da lockfile; per integrazione/HTTP/browser PostgreSQL locale isolato. Configurazione del runner.',
                         result['command'] + '; caso nominato nel log: ' + label,
                         expected,
                         observed + ' Exit code ' + str(result['exitCode']) + '.',
                         result['result'], result['startedAt'][:10], result['log'],
                         'Provider controllati nei test di integrazione/HTTP; non è una transazione remota.'])
    pending = [
        ('BUR-009, BUR-010 / FUD-002', 'Stripe TEST remoto', 'Eseguire acquisto TEST, ritorno e webhook; verificare un solo ordine pagato e importi coerenti.', 'BLOCKED', 'Credenziali provider e sessione browser anti-bot dell’ambiente remoto non disponibili nel collaudo locale.'),
        ('BUR-010, BUR-019 / FUD-008', 'Recapito email', 'Inviare una conferma a destinatario autorizzato e verificare ricezione e identificativo provider.', 'BLOCKED', 'Non eseguito invio esterno; nessuna credenziale Resend configurata nella verifica.'),
        ('BUR-013 / FUD-004', 'Rimborso remoto', 'Rimborsare un pagamento TEST e riconciliare stato, importo e notifiche.', 'BLOCKED', 'Dipende dalla transazione TEST remota e dalle credenziali autorizzate.'),
        ('BUR-022 / FUD-009', 'Turnstile e Redis remoti', 'Completare CAPTCHA sul dominio definitivo e verificare limite distribuito fra istanze.', 'BLOCKED', 'Token interattivo e credenziali remote non configurati.'),
        ('BUR-013, BUR-019 / FUD-004', 'Cron Vercel', 'Osservare invocazione pianificata, Bearer, log e smaltimento del backlog.', 'BLOCKED', 'Nessun deploy o cambio configurazione remoto del candidato eseguito.'),
        ('BUR-023 / FUD-010', 'Backup commerciale', 'Ripristinare snapshot attuale su destinazione separata; verificare dati e RPO/RTO.', 'BLOCKED', 'Eseguito soltanto backup/ripristino PostgreSQL locale; evidenze storiche non riutilizzate come nuova prova.'),
        ('BUR-024 / FUD-009', 'Dati venditore e condizioni', 'Validare dati autentici, contatti, condizioni e procedura di recesso del titolare.', 'BLOCKED', 'Dati e approvazione del titolare non forniti per il candidato.'),
        ('BUR-025 / FUD-001', 'Safari/iOS e Android fisici', 'Ripetere galleria, touch, zoom, CAPTCHA e checkout sui dispositivi reali.', 'NOT RUN', 'Mobile emulato Chromium verificato; hardware e browser differenti non eseguiti.'),
        ('BUR-026 / FUD-010', 'Carico e capacità', 'Misurare latenza, pool, contesa e backlog sul volume concordato.', 'NOT RUN', 'Non esiste un volume obiettivo concordato; nessuna capacità prestazionale dichiarata.'),
        ('BUR-015, BUR-016 / FUD-005', 'Admin del servizio', 'Applicare migrazione, inizializzare account autorizzato, login e logout sul dominio di destinazione.', 'BLOCKED', 'Creati esclusivamente account di collaudo isolati; credenziali reali non fornite.'),
        ('BUR-026 / FUD-010', 'Deploy candidato', 'Distribuire il commit approvato, verificare asset GLB, migrazioni, API e log.', 'NOT RUN', 'Branch locale preparato; nessuna pubblicazione o merge automatico.'),
    ]
    for req, module, procedure, status, reason in pending:
        rows.append([f'TC-{len(rows)+1:03}', req, module, module, 'Ambiente dedicato e accessi autorizzati.', procedure,
                     'Flusso completato con evidenza provider/dispositivo.', 'NON VERIFICATA: ' + reason, status, '', '', reason])
    return rows


def workbook(rows):
    wb = Workbook()
    sheet = wb.active
    sheet.title = 'Casi di test'
    headers = ['ID Test', 'Requisito', 'Modulo', 'Descrizione', 'Precondizioni', 'Procedura', 'Risultato atteso', 'Risultato effettivo', 'Esito', 'Data esecuzione', 'Evidenza', 'Note']
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    widths = [13, 36, 22, 52, 45, 65, 44, 56, 14, 19, 38, 55]
    for i, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(i)].width = width
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor=RED)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
        sheet.row_dimensions[row[0].row].height = 100
    table = Table(displayName='RegistroRUSH', ref=f'A1:L{sheet.max_row}')
    table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    sheet.add_table(table)
    sheet.freeze_panes = 'D2'
    sheet.print_title_rows = '1:1'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.oddFooter.center.text = 'RUSH Presale — Pagina &P / &N'
    validation = DataValidation(type='list', formula1='"PASS,FAIL,BLOCKED,NOT RUN"')
    sheet.add_data_validation(validation)
    validation.add(f'I2:I{sheet.max_row}')
    for status, color in [('PASS', 'D9EAD3'), ('FAIL', 'F4CCCC'), ('BLOCKED', 'FCE5CD'), ('NOT RUN', 'E7E6E6')]:
        sheet.conditional_formatting.add(f'I2:I{sheet.max_row}', CellIsRule(operator='equal', formula=['"' + status + '"'], fill=PatternFill('solid', fgColor=color)))
    summary = wb.create_sheet('Riepilogo')
    summary.append(['RUSH Presale — collaudo ' + VERSION, DATE])
    summary.append(['Stato', 'Numero casi'])
    for status in ['PASS', 'FAIL', 'BLOCKED', 'NOT RUN']:
        summary.append([status, f'=COUNTIF(\'Casi di test\'!I2:I{sheet.max_row},A{summary.max_row+1})'])
    summary.append(['Totale', f'=COUNTA(\'Casi di test\'!A2:A{sheet.max_row})'])
    summary.append(['Percentuale PASS sul totale', '=IFERROR(B3/B7,0)'])
    summary['B8'].number_format = '0.0%'
    summary.append(['Limite', 'PASS locale non certifica servizi remoti; BLOCKED e NOT RUN restano NON VERIFICATI.'])
    summary.column_dimensions['A'].width = 47
    summary.column_dimensions['B'].width = 95
    for row in summary:
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    guide = wb.create_sheet('Istruzioni')
    for row in [
        ['Versione documento', '1.0'], ['Data', DATE], ['Fonte', 'docs/evidenze/esecuzione.json e log conservati'],
        ['PASS', 'Asserzioni realmente eseguite, log e exit code disponibili.'],
        ['FAIL', 'Caso eseguito con esito negativo. I difetti corretti sono nel rapporto di collaudo.'],
        ['BLOCKED', 'Prerequisito mancante; data esecuzione vuota.'], ['NOT RUN', 'Prova non eseguita; nessun risultato attribuito.'],
        ['Tracciabilità', 'ID BUR e FUD presenti nei documenti corrispondenti; nomi inglesi dei test conservati per ritrovarli nei log.'],
        ['Aggiornamento', 'Registrare data reale, atteso/osservato ed evidenza senza segreti; le formule contano la colonna Esito.'],
        ['Ambiente', 'PostgreSQL locale isolato, Chromium desktop/mobile emulato, provider controllati.'],
    ]:
        guide.append(row)
    guide.column_dimensions['A'].width = 30
    guide.column_dimensions['B'].width = 120
    for row in guide:
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    path = DOCS / 'Casi_di_Test_RUSH_Presale.xlsx'
    wb.save(path)
    return path


def report(data, rows):
    text = '# Rapporto di collaudo RUSH Presale\n\n## 1. Identificazione e ambiente\n'
    text += f'Candidato {VERSION}, branch feature/rush-final-release, baseline 0e3ee36. Data delle prove: {DATE}. '
    text += data['scope'] + f' Runtime Node {data["node"]}; browser {data["browser"]["browser"]}. I log e gli hash delle esecuzioni sono conservati in docs/evidenze/esecuzione.json.\n'
    text += '\n## 2. Verifiche realmente eseguite\n\n| Comando | Esito | Evidenza |\n|---|---|---|\n'
    for row in data['results']:
        text += f'| {row["command"]} | {row["result"]}, exit {row["exitCode"]} | {row["log"]} |\n'
    text += '\nIl comando unitario esclude esplicitamente il contenitore PostgreSQL, eseguito separatamente da test:integration. I risultati del registro sono casi nominati e controlli di tool: non vanno sommati come se fossero tutti test unitari.\n'
    text += '\n## 3. Risultati per livello\n'
    for name in ['test', 'test-integration', 'test-http', 'test-browser']:
        count = len([row for row in rows if row[2] == name and row[8] == 'PASS'])
        text += f'{name}: {count} casi nominati PASS nel registro.\n\n'
    text += 'Gli handler HTTP sono stati verificati su Next compilato con PostgreSQL reale e webhook firmati tramite SDK su fixture. Stripe e Resend esterni non sono stati chiamati. Il browser usa il GLB originale e asset locali; le prove admin usano un account temporaneo con sessione DB valida.\n'
    text += '\n## 4. Anomalie riscontrate e correzioni\n\n| Anomalia | Correzione | Regresso verificato |\n|---|---|---|\n'
    text += '| Identità admin solo da ambiente | AdminUser persistente, ruolo ADMIN, bootstrap idempotente e controllo account su ogni sessione | Integrazione e HTTP |\n'
    text += '| CSP bloccava WebAssembly | Aggiunta wasm-unsafe-eval, senza abilitare eval JavaScript | Browser da home a PRE-SALE e test CSP |\n'
    text += '| Controllo HTTP confondeva wasm-unsafe-eval con unsafe-eval | Confronto dei token CSP completi; eval JavaScript resta vietato | Suite HTTP ripetuta dopo la correzione del test |\n'
    text += '| Texture GLB non renderizzate per blocco dei blob | connect-src consente blob locali; controllo errori console oltre alle eccezioni | Texture e screenshot browser |\n'
    text += '| Inquadratura originale tagliava il modello | Distanza iniziale adattiva 105%, limiti zoom espliciti | Screenshot desktop/mobile e ripristino |\n'
    text += '| Home con reduced-motion causava mismatch di idratazione | Markup particelle stabile e disattivazione tramite media query CSS | Browser reduced-motion, console senza errori |\n'
    text += '| Assenza WebGL non produceva tempestivamente il fallback | Controllo WebGL2 prima di importare il motore | Browser avviato con WebGL disabilitato |\n'
    text += '| Reinstallazione bloccata su esbuild in uso dal runner | Verifica release avviata con Node nativo senza tsx | npm ci e sequenza finale completati |\n'
    text += '| Sandbox Windows impediva avvio tsx e PostgreSQL | Esecuzione dei runner autorizzati fuori sandbox, sempre su loopback isolato | Esecuzioni finali registrate |\n'
    text += '\nUn tentativo browser ha incontrato un timeout del renderer software; la ripetizione con il renderer standard ha prodotto le evidenze finali. I tentativi intermedi non sono usati come PASS. Nessun dato di produzione è stato modificato.\n'
    text += '\n## 5. Screenshot e verifiche visive\n'
    text += 'Le immagini del manuale provengono dal browser reale sul server locale: home, prodotto, viewer fronte/retro, mobile, fallback, login, dashboard e impostazioni. Il mobile è emulato in Chromium, non un certificato Safari/iOS. È verificata la presenza delle texture e l’assenza di errori CSP nel percorso positivo.\n'
    text += '\n## 6. Prove non verificate\n\n| Modulo | Stato | Causa |\n|---|---|---|\n'
    for row in rows:
        if row[8] != 'PASS':
            text += f'| {row[2]} | {row[8]} — NON VERIFICATA | {row[11]} |\n'
    text += '\n## 7. Pulizia e sicurezza della consegna\n'
    text += 'I cinque documenti DOCX e il registro XLSX sono verificati per integrità ZIP/XML, riapertura, struttura e formule. Il rendering in Microsoft Office non è stato eseguito. Il registro comprende casi nominati e controlli degli strumenti; i riepiloghi conservano separatamente PASS, FAIL, BLOCKED e NOT RUN.\n\n'
    text += 'I test automatici e la diagnostica admin sono conservati. La pagina /anteprima resta un redirect di compatibilità; il login staging ritirato restituisce 410. Nessuna pagina demo del pacchetto 3D è pubblicata. Il banner TEST è mantenuto negli ambienti di collaudo e assente con production/live, senza nascondere la modalità dei pagamenti. Le tabelle storiche ignorate restano intatte.\n'
    text += '\n## 8. Valutazione finale del rilascio\n'
    text += 'Il candidato ha superato le verifiche tecniche locali registrate ed è preparato per revisione tramite Pull Request. Il rilascio commerciale LIVE non è certificato: richiede bootstrap dell’admin reale, migrazione sulla destinazione verificata, collaudo provider, dati venditore, backup e scheduler del servizio. Non è stato eseguito un deploy né un merge del candidato.\n'
    source = SOURCES / 'Rapporto_Collaudo_RUSH_Presale.md'
    source.write_text(text, encoding='utf-8')
    return source


if __name__ == '__main__':
    if '--capture' in sys.argv:
        capture()
    data = json.loads((EVIDENCE / 'esecuzione.json').read_text(encoding='utf-8'))
    DATE = datetime.fromisoformat(data['date']).strftime('%d/%m/%Y')
    diagrams()
    rows = cases(data)
    report(data, rows)
    paths = []
    for source in sorted(SOURCES.glob('*.md')):
        paths.append(make_docx(source, technical_appendix() if source.stem.startswith('TAD_') else ''))
    paths.append(workbook(rows))
    for path in paths:
        with zipfile.ZipFile(path) as archive:
            assert archive.testzip() is None
            for name in archive.namelist():
                if name.endswith(('.xml', '.rels')):
                    ET.fromstring(archive.read(name))
        if path.suffix == '.docx':
            loaded = Document(path)
            assert len(loaded.paragraphs) > 20 and loaded.sections[0].footer.paragraphs
        else:
            loaded = load_workbook(path)
            assert loaded['Casi di test'].max_row == len(rows) + 1
            assert loaded['Riepilogo']['B3'].data_type == 'f'
        print('VALIDATO ' + path.name)
    (EVIDENCE / 'documenti.json').write_text(json.dumps({
        'validation': 'ZIP, XML, riapertura DOCX/XLSX, struttura e formule; rendering Microsoft Office non eseguito.',
        'date': DATE, 'cases': len(rows), 'statuses': {status: sum(row[8] == status for row in rows) for status in ['PASS', 'FAIL', 'BLOCKED', 'NOT RUN']},
        'files': [{'path': path.name, 'bytes': path.stat().st_size, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest()} for path in paths]
    }, ensure_ascii=False, indent=2), encoding='utf-8')
